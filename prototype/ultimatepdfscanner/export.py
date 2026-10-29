"""Exporters: searchable PDF (image + invisible OCR text layer), TXT, DOCX."""
import io
from pathlib import Path

import cv2
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from .types import Page

FONT = "Helvetica"


def _pil(image) -> Image.Image:
    if image.ndim == 2:
        return Image.fromarray(image)
    return Image.fromarray(cv2.cvtColor(image, cv2.COLOR_BGR2RGB))


def _encoded(image) -> ImageReader:
    pil = _pil(image)
    buf = io.BytesIO()
    if pil.mode == "L" and set(pil.getextrema()) <= {0, 255}:
        pil.convert("1").save(buf, "PNG")  # bilevel scans compress far better losslessly
    else:
        pil.save(buf, "JPEG", quality=85)
    buf.seek(0)
    return ImageReader(buf)


def _latin1(text: str) -> str:
    return text.encode("latin-1", errors="replace").decode("latin-1")


def export_pdf(pages: list[Page], path: Path, dpi: int = 300) -> None:
    pt = 72.0 / dpi
    c = canvas.Canvas(str(path))
    c.setTitle(Path(path).stem)
    for page in pages:
        h_px, w_px = page.image.shape[:2]
        pw, ph = w_px * pt, h_px * pt
        c.setPageSize((pw, ph))
        c.drawImage(_encoded(page.image), 0, 0, pw, ph)
        for w in page.words:
            x0, y0, x1, y1 = w.box
            size = max(1.0, (y1 - y0) * pt)
            text = _latin1(w.text)
            natural = c.stringWidth(text, FONT, size)
            t = c.beginText()
            t.setTextRenderMode(3)  # invisible: selectable/searchable but not drawn
            t.setFont(FONT, size)
            if natural > 0:
                t.setHorizScale(100.0 * (x1 - x0) * pt / natural)
            t.setTextOrigin(x0 * pt, ph - y1 * pt + 0.2 * size)
            t.textOut(text)
            c.drawText(t)
        c.showPage()
    c.save()


def export_txt(pages: list[Page], path: Path) -> None:
    Path(path).write_text("\n\f\n".join(p.text for p in pages), encoding="utf-8")


def export_docx(pages: list[Page], path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document()
    for i, page in enumerate(pages):
        if i:
            doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        for para in page.text.split("\n\n") or [""]:
            doc.add_paragraph(para)
    doc.save(str(path))
